"""Excel / CSV import for the guided pipeline.

The flow has three server calls, all stateless (the file is re-sent each time):

1. parse the file -> return its column headers for the mapping screen;
2. preview -- apply a column->field mapping, resolve foreign keys by their
   human-readable natural key, validate every row, and report per-row flags;
3. commit -- the same, but actually insert the rows that passed.

Foreign keys are resolved by natural key (tag / code / model number), never by
raw id, and an unresolved reference flags only that row -- the rest still import.
"""

from __future__ import annotations

import csv
import enum
import io
import typing

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from .registry import TableSpec, fk_target_table_name, spec_for_table_name

# Columns never imported (assigned by the server / database).
_SKIP_FIELDS = {"project_id", "id", "created_at", "updated_at"}


# --------------------------------------------------------------------------
# Spreadsheet parsing
# --------------------------------------------------------------------------


def parse_spreadsheet(filename: str, content: bytes) -> tuple[list[str], list[dict]]:
    """Parse an .xlsx or .csv file into (headers, rows-as-dicts)."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _parse_csv(content)
    if name.endswith((".xlsx", ".xlsm")):
        return _parse_xlsx(content)
    raise ValueError("Unsupported file type. Upload a .xlsx or .csv file.")


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            return [], []
        headers = [
            str(h).strip() if h is not None else f"Column{i + 1}"
            for i, h in enumerate(header)
        ]
        data: list[dict] = []
        for row in rows_iter:
            if row is None or all(v is None for v in row):
                continue
            data.append(
                {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}
            )
        return headers, data
    finally:
        workbook.close()


def _parse_csv(content: bytes) -> tuple[list[str], list[dict]]:
    text = content.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return [], []
    headers = [h.strip() for h in rows[0]]
    data: list[dict] = []
    for row in rows[1:]:
        if not any((cell or "").strip() for cell in row):
            continue
        data.append(
            {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}
        )
    return headers, data


# --------------------------------------------------------------------------
# Field metadata
# --------------------------------------------------------------------------


def importable_fields(spec: TableSpec) -> list[str]:
    """Fields a user may supply when importing rows for this table."""
    return [f for f in spec.create_schema.model_fields if f not in _SKIP_FIELDS]


def _clean(value: object) -> object | None:
    """Normalise a cell value; blank strings become None."""
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value


# --------------------------------------------------------------------------
# Foreign-key resolution
# --------------------------------------------------------------------------


def _resolve_fk(
    db: Session, target: TableSpec, project_id: int, value: object
) -> int | None:
    """Resolve a foreign key by the target table's natural key."""
    key_column = getattr(target.model, target.natural_key)
    stmt = select(target.model.id).where(key_column == str(value).strip())
    if target.project_scoped:
        stmt = stmt.where(target.model.project_id == project_id)
    return db.scalar(stmt)


def _process_row(
    db: Session,
    spec: TableSpec,
    project_id: int,
    raw: dict,
    mapping: dict[str, str],
    line_number: int,
) -> dict:
    """Map, resolve and validate one spreadsheet row."""
    values: dict[str, object] = {}
    errors: list[str] = []

    for column, field in mapping.items():
        if not field or column not in raw:
            continue
        value = _clean(raw.get(column))
        target_table = fk_target_table_name(spec.model, field)

        if target_table and value is not None:
            target_spec = spec_for_table_name(target_table)
            if target_spec is None:
                errors.append(f"{field}: cannot resolve references to {target_table}.")
                continue
            resolved = _resolve_fk(db, target_spec, project_id, value)
            if resolved is None:
                errors.append(
                    f"{field}: '{value}' was not found in {target_spec.label} "
                    f"(matched on {target_spec.natural_key})."
                )
            else:
                values[field] = resolved
        else:
            values[field] = value

    if spec.project_scoped:
        values["project_id"] = project_id

    # Validate against the same schema a normal create uses.
    try:
        spec.create_schema(**values)
    except ValidationError as exc:
        for err in exc.errors():
            loc = ".".join(str(x) for x in err.get("loc", ()))
            errors.append(f"{loc or 'row'}: {err.get('msg')}")

    return {
        "row": line_number,
        "values": {k: v for k, v in values.items() if k != "project_id"},
        "errors": errors,
        "status": "error" if errors else "ok",
    }


# --------------------------------------------------------------------------
# Preview & commit
# --------------------------------------------------------------------------


def preview_import(
    db: Session,
    spec: TableSpec,
    project_id: int,
    rows: list[dict],
    mapping: dict[str, str],
) -> dict:
    """Validate every row without inserting anything."""
    results = [
        _process_row(db, spec, project_id, raw, mapping, i)
        for i, raw in enumerate(rows, start=2)  # row 1 is the header
    ]
    ok = sum(1 for r in results if r["status"] == "ok")
    return {
        "total": len(results),
        "ok_count": ok,
        "error_count": len(results) - ok,
        "rows": results[:1000],
    }


def commit_import(
    db: Session,
    spec: TableSpec,
    project_id: int,
    rows: list[dict],
    mapping: dict[str, str],
) -> dict:
    """Insert every row that passes validation; report per-row failures.

    Rows are processed in order and committed one at a time, so a row may
    reference an earlier row of the same import (e.g. a self-referencing tag).
    """
    created = 0
    errors: list[dict] = []
    for i, raw in enumerate(rows, start=2):
        result = _process_row(db, spec, project_id, raw, mapping, i)
        if result["status"] != "ok":
            errors.append({"row": result["row"], "errors": result["errors"]})
            continue
        try:
            full_values = dict(result["values"])
            if spec.project_scoped:
                full_values["project_id"] = project_id
            db.add(spec.model(**full_values))
            db.commit()
            created += 1
        except Exception as exc:  # noqa: BLE001 -- reported per row
            db.rollback()
            errors.append(
                {"row": result["row"], "errors": [str(getattr(exc, "orig", exc))]}
            )
    return {"created": created, "total": len(rows), "errors": errors[:500]}


# --------------------------------------------------------------------------
# Template generation
# --------------------------------------------------------------------------


def _enum_example(annotation: object) -> str | None:
    """First value of an enum annotation (unwrapping Optional), if any."""
    for candidate in (annotation, *typing.get_args(annotation)):
        if isinstance(candidate, type) and issubclass(candidate, enum.Enum):
            return list(candidate)[0].value
    return None


def _example_value(spec: TableSpec, field: str) -> object:
    """A generic placeholder value for the template's example row."""
    info = spec.create_schema.model_fields[field]
    if field == spec.natural_key:
        return "EXAMPLE-1"
    enum_value = _enum_example(info.annotation)
    if enum_value is not None:
        return enum_value
    if fk_target_table_name(spec.model, field):
        return ""  # FK: leave blank, fill with a tag/code when used
    return ""


def build_template(spec: TableSpec) -> bytes:
    """Build an .xlsx template: a header row plus one example row."""
    fields = importable_fields(spec)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = spec.slug.replace("project-", "")[:31]
    sheet.append(fields)
    sheet.append([_example_value(spec, f) for f in fields])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
